from __future__ import annotations

import logging
import re
import shutil
import textwrap
from pathlib import Path
from datetime import datetime
from typing import Iterable, Iterator, List, Sequence, Tuple

import pdfplumber
from docx import Document as DocxDocument
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy.orm import Session

from .config import settings
from .db.models import Document, Chunk, Page, Region
from .ocr_utils import ocr_image_fast, ocr_image_tiled
from .vectorstore import VectorStores

logger = logging.getLogger(__name__)

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

_WHITESPACE_RX = re.compile(r"[ \t\u00a0\u202f]+")
_MULTILINE_RX = re.compile(r"\n{3,}")
_SENTENCE_SPLIT_RX = re.compile(r"(?<=[\.!?…])\s+(?=[A-ZА-ЯЁ0-9])")
_PAGE_BREAK_TOKEN = "[[PAGE_BREAK]]"
_DOCX_PAGE_TARGET_CHARS = 3200
_PREVIEW_PAGE_SIZE = (1240, 1754)  # A4-like preview
_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
]


def normalize_text(text: str) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = []
    for raw in text.split("\n"):
        line = _WHITESPACE_RX.sub(" ", raw).strip()
        lines.append(line)
    text = "\n".join(lines)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = _MULTILINE_RX.sub("\n\n", text)
    return text.strip()


def _sentence_chunks(paragraph: str, *, target_size: int, overlap_chars: int) -> List[str]:
    paragraph = normalize_text(paragraph)
    if not paragraph:
        return []
    if len(paragraph) <= target_size:
        return [paragraph]

    sentences = [s.strip() for s in _SENTENCE_SPLIT_RX.split(paragraph) if s.strip()]
    if len(sentences) <= 1:
        step = max(1, target_size - overlap_chars)
        return [paragraph[i:i + target_size] for i in range(0, len(paragraph), step)]

    out: List[str] = []
    current: List[str] = []
    current_len = 0
    for sent in sentences:
        extra = len(sent) + (1 if current else 0)
        if current and current_len + extra > target_size:
            out.append(" ".join(current).strip())
            overlap: List[str] = []
            overlap_len = 0
            for prev in reversed(current):
                overlap.insert(0, prev)
                overlap_len += len(prev) + 1
                if overlap_len >= overlap_chars:
                    break
            current = overlap[:]
            current_len = sum(len(x) for x in current) + max(0, len(current) - 1)
        current.append(sent)
        current_len += extra
    if current:
        out.append(" ".join(current).strip())
    return [x for x in out if x]


def _table_chunks(paragraph: str, *, target_size: int, overlap_chars: int) -> List[str]:
    lines = [line.strip() for line in paragraph.split("\n") if line.strip()]
    if len(paragraph) <= target_size or len(lines) <= 2:
        return [paragraph]

    title = lines[0]
    chunks: List[str] = []
    current: List[str] = [title]

    for line in lines[1:]:
        candidate = "\n".join([*current, line])
        if len(current) > 1 and len(candidate) > target_size:
            chunks.append("\n".join(current).strip())
            overlap_rows: List[str] = []
            overlap_len = 0
            for prev in reversed(current[1:]):
                overlap_rows.insert(0, prev)
                overlap_len += len(prev) + 1
                if overlap_len >= overlap_chars:
                    break
            current = [title, *overlap_rows]
        current.append(line)

    if len(current) > 1:
        chunks.append("\n".join(current).strip())
    return [chunk for chunk in chunks if chunk]


def chunk_text(text: str, *, chunk_size: int, overlap: int) -> List[str]:
    text = normalize_text(text)
    if not text:
        return []
    if chunk_size <= 0:
        return [text]

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: List[str] = []
    current: List[str] = []
    current_len = 0

    for paragraph in paragraphs:
        if paragraph.startswith("[TABLE ") and len(paragraph) > chunk_size:
            if current:
                chunks.append("\n\n".join(current).strip())
                current = []
                current_len = 0
            chunks.extend(_table_chunks(paragraph, target_size=chunk_size, overlap_chars=overlap))
            continue

        if len(paragraph) > chunk_size:
            if current:
                chunks.append("\n\n".join(current).strip())
                current = []
                current_len = 0
            chunks.extend(_sentence_chunks(paragraph, target_size=chunk_size, overlap_chars=overlap))
            continue

        addition = len(paragraph) + (2 if current else 0)
        if current and current_len + addition > chunk_size:
            chunks.append("\n\n".join(current).strip())
            overlap_text = ""
            joined = "\n\n".join(current).strip()
            if overlap > 0 and joined:
                overlap_text = joined[-overlap:].strip()
            current = [overlap_text, paragraph] if overlap_text else [paragraph]
            current = [x for x in current if x]
            current_len = len("\n\n".join(current))
        else:
            current.append(paragraph)
            current_len += addition

    if current:
        chunks.append("\n\n".join(current).strip())

    min_chars = max(1, int(getattr(settings, "CHUNK_MIN_CHARS", 120)))
    merged: List[str] = []
    for ch in chunks:
        if merged and len(ch) < min_chars:
            merged[-1] = f"{merged[-1]}\n\n{ch}".strip()
        else:
            merged.append(ch)
    return [normalize_text(ch) for ch in merged if normalize_text(ch)]


def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in _FONT_CANDIDATES:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size=size)
            except Exception:
                continue
    return ImageFont.load_default()


def _save_docx_preview(*, page_text: str, out_path: Path, filename: str, page_number: int) -> tuple[int, int]:
    width, height = _PREVIEW_PAGE_SIZE
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _load_font(26)
    body_font = _load_font(22)
    small_font = _load_font(18)

    margin_x = 72
    y = 56
    draw.text((margin_x, y), f"{filename} — лист {page_number}", fill=(30, 30, 30), font=title_font)
    y += 46
    draw.line((margin_x, y, width - margin_x, y), fill=(210, 210, 210), width=2)
    y += 26

    max_width = width - margin_x * 2
    avg_char_px = max(8, body_font.size // 2)
    wrap_chars = max(40, max_width // avg_char_px)
    wrapped_lines: List[str] = []
    for paragraph in normalize_text(page_text).split("\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            wrapped_lines.append("")
            continue
        wrapped_lines.extend(textwrap.wrap(paragraph, width=wrap_chars, break_long_words=False, break_on_hyphens=False) or [paragraph])
        wrapped_lines.append("")

    line_height = int(body_font.size * 1.5)
    max_lines = max(10, (height - y - 120) // line_height)
    clipped = False
    if len(wrapped_lines) > max_lines:
        wrapped_lines = wrapped_lines[: max_lines - 2] + ["", "[Текст сокращён в превью. Для копирования используйте скачивание оригинала DOCX.]"]
        clipped = True

    for line in wrapped_lines:
        draw.text((margin_x, y), line, fill=(20, 20, 20), font=small_font if line.startswith("[Текст сокращён") else body_font)
        y += line_height
        if y >= height - 80:
            clipped = True
            break

    footer = "Текстовое превью страницы DOCX" + (" • часть текста сокращена" if clipped else "")
    draw.text((margin_x, height - 54), footer, fill=(100, 100, 100), font=small_font)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path)
    image.close()
    return width, height


def _require_fitz() -> None:
    if fitz is None:
        raise RuntimeError("PyMuPDF is required for PDF processing")


def _extract_pdf_text_with_pdfplumber(path: Path, page_index: int) -> str:
    try:
        with pdfplumber.open(str(path)) as pdf:
            page = pdf.pages[page_index]
            return normalize_text(page.extract_text(x_tolerance=2, y_tolerance=3, layout=True) or page.extract_text() or "")
    except Exception:
        return ""


def _table_rows_to_text(rows: Sequence[Sequence[object | None]]) -> str:
    lines: List[str] = []
    for row in rows or []:
        cells = [normalize_text(str(cell or "")) for cell in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            lines.append(" | ".join(cell or "-" for cell in cells))
    return "\n".join(lines).strip()


def _extract_pdf_tables_with_pdfplumber(path: Path, page_index: int) -> List[str]:
    table_settings = {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 5,
    }
    try:
        with pdfplumber.open(str(path)) as pdf:
            page = pdf.pages[page_index]
            tables = page.extract_tables(table_settings=table_settings) or []
            if not tables:
                tables = page.extract_tables() or []
    except Exception as exc:
        logger.warning("PDF table extraction failed: file=%s page=%s: %s", path, page_index + 1, exc)
        return []

    blocks: List[str] = []
    for table_index, table in enumerate(tables, start=1):
        table_text = _table_rows_to_text(table)
        if table_text:
            blocks.append(f"[TABLE {table_index}]\n{table_text}")
    return blocks


def _is_text_pdf(text: str, words_count: int) -> bool:
    cleaned = re.sub(r"\s+", " ", text or "").strip()
    alnum = re.sub(r"[^\wА-Яа-яЁё]+", "", cleaned, flags=re.UNICODE)
    return len(alnum) >= 40 or words_count >= 18 or len(cleaned) >= 120


def _ocr_page_image(img: Image.Image) -> str:
    lang = getattr(settings, "OCR_LANG", "rus+eng")
    timeout = int(getattr(settings, "OCR_TIMEOUT_SEC", 20))
    max_side = int(getattr(settings, "OCR_FAST_MAX_SIDE", 1600))
    try:
        o = ocr_image_tiled(img, lang=lang, timeout_sec=timeout, max_side=max_side, tiles_x=2, tiles_y=2)
        text = normalize_text(o.text)
        if text:
            return text
    except Exception as exc:
        logger.warning("OCR tiled failed, falling back to fast OCR: %s", exc)
    try:
        o = ocr_image_fast(img, lang=lang, timeout_sec=timeout, max_side=max_side)
        return normalize_text(o.text)
    except Exception as exc:
        logger.warning("OCR fast failed: %s", exc)
        return ""


def _render_pdf_page_image(page: "fitz.Page", dpi: int) -> Image.Image:
    scale = dpi / 72.0
    max_side = int(getattr(settings, "RENDER_MAX_SIDE", 0) or 0)
    if max_side > 0:
        try:
            rect = page.rect
            rendered_w = float(rect.width) * scale
            rendered_h = float(rect.height) * scale
            largest = max(rendered_w, rendered_h)
            if largest > max_side:
                scale *= max_side / largest
        except Exception:
            pass
    scale = max(0.2, scale)
    mat = fitz.Matrix(scale, scale)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    return Image.frombytes("RGB", [pix.width, pix.height], pix.samples)


def _extract_pdf_pages(path: Path, *, preview_dir: Path, filename: str) -> List[dict]:
    _require_fitz()
    pdf = fitz.open(str(path))
    pages: List[dict] = []
    try:
        max_pages = int(getattr(settings, "MAX_PAGES", 0) or 0)
        page_count = min(int(pdf.page_count), max_pages) if max_pages > 0 else int(pdf.page_count)
        if max_pages > 0 and int(pdf.page_count) > page_count:
            logger.warning(
                "PDF page processing capped: file=%s total_pages=%s max_pages=%s",
                filename,
                int(pdf.page_count),
                page_count,
            )
        for page_index in range(page_count):
            page = pdf.load_page(page_index)
            page_number = page_index + 1
            if page_number == 1 or page_number % 10 == 0 or page_number == page_count:
                logger.info("Extracting PDF page: file=%s page=%s/%s", filename, page_number, page_count)
            preview_path = preview_dir / f"page_{page_number:04d}.png"
            img = _render_pdf_page_image(page, dpi=int(getattr(settings, "RENDER_DPI", 220)))
            img.save(preview_path, compress_level=3)

            text_layer = normalize_text(page.get_text("text", sort=True) or "")
            words_count = len(page.get_text("words") or [])
            if not _is_text_pdf(text_layer, words_count):
                plumber_text = _extract_pdf_text_with_pdfplumber(path, page_index)
                if len(plumber_text) > len(text_layer):
                    text_layer = plumber_text
            is_scan = not _is_text_pdf(text_layer, words_count)
            page_text = text_layer
            source_kind = "pdf_text"
            if is_scan:
                page_text = _ocr_page_image(img)
                source_kind = "ocr"
            if not page_text:
                page_text = text_layer or ""
            page_text = normalize_text(page_text)
            if bool(getattr(settings, "ENABLE_TABLE_DETECTION", False)):
                table_blocks = _extract_pdf_tables_with_pdfplumber(path, page_index)
                if table_blocks:
                    page_text = normalize_text(f"{page_text}\n\n" + "\n\n".join(table_blocks))
                    source_kind = f"{source_kind}+tables"
            pages.append(
                {
                    "page_number": page_number,
                    "text": page_text,
                    "width": img.width,
                    "height": img.height,
                    "image_path": str(preview_path),
                    "source_kind": source_kind,
                }
            )
            img.close()
    finally:
        pdf.close()
    logger.info("Extracted PDF pages: file=%s pages=%s", filename, len(pages))
    return pages


def _iter_block_items(parent: _Document | _Cell) -> Iterator[Paragraph | Table]:
    parent_elm = parent.element.body if isinstance(parent, _Document) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _table_to_text(table: Table) -> str:
    rows: List[str] = []
    for row in table.rows:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(c or "—" for c in cells))
    return "\n".join(rows).strip()


def _paragraph_has_page_break(paragraph: Paragraph) -> bool:
    for run in paragraph.runs:
        xml = getattr(getattr(run, "_r", None), "xml", "") or ""
        if 'w:type="page"' in xml or "w:type='page'" in xml:
            return True
    return False


def _docx_header_footer_text(doc: DocxDocument) -> str:
    parts: List[str] = []
    for section in doc.sections:
        header_text = normalize_text("\n".join(p.text for p in section.header.paragraphs if p.text.strip()))
        footer_text = normalize_text("\n".join(p.text for p in section.footer.paragraphs if p.text.strip()))
        if header_text:
            parts.append(f"[HEADER]\n{header_text}")
        if footer_text:
            parts.append(f"[FOOTER]\n{footer_text}")
    return normalize_text("\n\n".join(parts))


def _extract_docx_pages(path: Path, *, preview_dir: Path, filename: str) -> List[dict]:
    doc = DocxDocument(str(path))
    header_footer = _docx_header_footer_text(doc)
    pages: List[dict] = []
    current_blocks: List[str] = []
    current_chars = 0

    def flush_page() -> None:
        nonlocal current_blocks, current_chars
        page_text = normalize_text("\n\n".join([b for b in current_blocks if b]).replace(_PAGE_BREAK_TOKEN, ""))
        if not page_text:
            current_blocks = []
            current_chars = 0
            return
        page_number = len(pages) + 1
        if page_number == 1 and header_footer:
            page_text = normalize_text(f"{header_footer}\n\n{page_text}")
        preview_path = preview_dir / f"page_{page_number:04d}.png"
        width, height = _save_docx_preview(page_text=page_text, out_path=preview_path, filename=filename, page_number=page_number)
        pages.append(
            {
                "page_number": page_number,
                "text": page_text,
                "width": width,
                "height": height,
                "image_path": str(preview_path),
                "source_kind": "docx_text",
            }
        )
        current_blocks = []
        current_chars = 0

    for block in _iter_block_items(doc):
        text = ""
        page_break_after = False
        if isinstance(block, Paragraph):
            text = normalize_text(block.text)
            page_break_after = _paragraph_has_page_break(block)
        else:
            table_text = _table_to_text(block)
            if table_text:
                text = f"[TABLE]\n{table_text}"
        if text:
            current_blocks.append(text)
            current_chars += len(text)
        if page_break_after or current_chars >= _DOCX_PAGE_TARGET_CHARS:
            flush_page()

    if current_blocks:
        flush_page()
    if not pages and header_footer:
        preview_path = preview_dir / "page_0001.png"
        width, height = _save_docx_preview(page_text=header_footer, out_path=preview_path, filename=filename, page_number=1)
        pages.append(
            {
                "page_number": 1,
                "text": header_footer,
                "width": width,
                "height": height,
                "image_path": str(preview_path),
                "source_kind": "docx_text",
            }
        )
    logger.info("Extracted DOCX pages: file=%s pages=%s", filename, len(pages))
    return pages


def _iter_chunks_with_pages(page_items: Sequence[dict]) -> Iterable[Tuple[int, str]]:
    for page in page_items:
        page_number = int(page.get("page_number") or 0)
        page_text = normalize_text(str(page.get("text") or ""))
        if page_number <= 0 or not page_text:
            continue
        for chunk in chunk_text(page_text, chunk_size=settings.CHUNK_SIZE, overlap=settings.CHUNK_OVERLAP):
            if chunk:
                yield page_number, chunk




def _set_document_state(
    db: Session,
    doc: Document,
    *,
    status: str | None = None,
    stage: str | None = None,
    progress: int | None = None,
    detail: str | None = None,
    finished: bool = False,
) -> None:
    if status is not None:
        doc.status = status
    if stage is not None:
        doc.stage = stage
    if progress is not None:
        try:
            doc.progress = max(0, min(100, int(progress)))
        except Exception:
            doc.progress = 0
    if detail is not None:
        doc.detail = str(detail)
    if getattr(doc, "processing_started_at", None) is None:
        doc.processing_started_at = datetime.utcnow()
    if finished:
        doc.processing_finished_at = datetime.utcnow()
    db.add(doc)
    db.commit()
    db.refresh(doc)


def _cleanup_failed_document(db: Session, doc: Document) -> None:
    stored_path = str(getattr(doc, "stored_path", "") or "")
    page_image_paths = [str(getattr(page, "image_path", "") or "") for page in list(getattr(doc, "pages", []) or [])]
    page_dir = None
    if page_image_paths:
        try:
            page_dir = str(Path(page_image_paths[0]).parent)
        except Exception:
            page_dir = None
    db.delete(doc)
    db.commit()
    for raw_path in [stored_path, *page_image_paths]:
        if not raw_path:
            continue
        try:
            Path(raw_path).unlink(missing_ok=True)
        except Exception:
            logger.warning("Failed to remove file during cleanup: %s", raw_path)
    if page_dir:
        try:
            Path(page_dir).rmdir()
        except Exception:
            pass


def _remove_tree(path: Path) -> None:
    try:
        if path.exists():
            shutil.rmtree(path, ignore_errors=True)
    except Exception:
        logger.warning("Failed to remove directory: %s", path)


def _replace_tree(src: Path, dst: Path) -> None:
    _remove_tree(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    if src.exists():
        shutil.move(str(src), str(dst))


def _rewrite_page_image_paths(page_items: Sequence[dict], *, src_dir: Path, dst_dir: Path) -> List[dict]:
    rewritten: List[dict] = []
    try:
        src_resolved = src_dir.resolve()
    except Exception:
        src_resolved = src_dir
    for item in page_items:
        row = dict(item)
        raw = str(row.get("image_path") or "")
        if raw:
            try:
                path = Path(raw)
                if path.parent.resolve() == src_resolved:
                    row["image_path"] = str(dst_dir / path.name)
            except Exception:
                pass
        rewritten.append(row)
    return rewritten


def _detect_file_type_from_path(path: Path, filename: str) -> str:
    name = f"{filename or ''} {path.name}".lower()
    if name.endswith(".pdf") or path.suffix.lower() == ".pdf":
        return "pdf"
    if name.endswith(".docx") or path.suffix.lower() == ".docx":
        return "docx"
    raise ValueError("Only PDF and DOCX documents can be reindexed")


def _delete_document_derivatives(
    *,
    db: Session,
    vectorstores: VectorStores,
    organization_id: int | None,
    project_id: int,
    document_id: int,
) -> tuple[int, int]:
    where = {"organization_id": int(organization_id or 0), "project_id": int(project_id), "document_id": int(document_id)}
    if bool(getattr(settings, "INDEX_EMBEDDINGS", False)):
        try:
            vectorstores.delete_text_where(where)
        except Exception:
            logger.warning("Failed to delete text embeddings before reindex: document_id=%s", document_id)
        try:
            vectorstores.delete_asset_where(where)
        except Exception:
            logger.warning("Failed to delete asset embeddings before reindex: document_id=%s", document_id)

    old_chunks = db.query(Chunk).filter(Chunk.document_id == int(document_id)).count()
    pages = db.query(Page).filter(Page.document_id == int(document_id)).all()
    page_ids = [int(page.id) for page in pages]
    old_regions = db.query(Region).filter(Region.page_id.in_(page_ids)).count() if page_ids else 0

    if page_ids:
        db.query(Region).filter(Region.page_id.in_(page_ids)).delete(synchronize_session=False)
    db.query(Page).filter(Page.document_id == int(document_id)).delete(synchronize_session=False)
    db.query(Chunk).filter(Chunk.document_id == int(document_id)).delete(synchronize_session=False)
    db.commit()
    return int(old_chunks), int(old_regions)


def _should_index_embeddings() -> bool:
    return bool(getattr(settings, "INDEX_EMBEDDINGS", False))


def reindex_document_text(
    *,
    db: Session,
    vectorstores: VectorStores,
    document_id: int,
    project_id: int,
    organization_id: int | None,
) -> Tuple[int, int, int, int]:
    doc = (
        db.query(Document)
        .filter(
            Document.id == int(document_id),
            Document.project_id == int(project_id),
            Document.organization_id == organization_id,
        )
        .first()
    )
    if not doc:
        raise ValueError(f"Document not found: project_id={project_id} document_id={document_id}")

    stored_path = Path(str(doc.stored_path or ""))
    if not stored_path.exists() or not stored_path.is_file():
        raise FileNotFoundError(f"Original file not found for document_id={document_id}")

    original_filename = str(doc.filename or stored_path.name)
    file_type = _detect_file_type_from_path(stored_path, original_filename)
    final_page_dir = settings.PAGES_DIR / f"org_{organization_id or 0}" / str(project_id) / str(doc.id)
    tmp_page_dir = settings.PAGES_DIR / f"org_{organization_id or 0}" / str(project_id) / f"{doc.id}_reindex_tmp"
    regions_dir = settings.REGIONS_DIR / f"org_{organization_id or 0}" / str(project_id) / str(doc.id)

    _remove_tree(tmp_page_dir)
    tmp_page_dir.mkdir(parents=True, exist_ok=True)

    _set_document_state(
        db,
        doc,
        status="processing",
        stage="Пересборка индекса",
        progress=10,
        detail="Исходный файл найден, запускаем повторное извлечение",
    )

    try:
        page_items = (
            _extract_pdf_pages(stored_path, preview_dir=tmp_page_dir, filename=original_filename)
            if file_type == "pdf"
            else _extract_docx_pages(stored_path, preview_dir=tmp_page_dir, filename=original_filename)
        )

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Очистка старого индекса",
            progress=45,
            detail=f"Повторно извлечено страниц: {len(page_items)}",
        )

        old_chunks, old_regions = _delete_document_derivatives(
            db=db,
            vectorstores=vectorstores,
            organization_id=organization_id,
            project_id=project_id,
            document_id=int(doc.id),
        )
        _remove_tree(regions_dir)
        _replace_tree(tmp_page_dir, final_page_dir)
        page_items = _rewrite_page_image_paths(page_items, src_dir=tmp_page_dir, dst_dir=final_page_dir)

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Формирование нового индекса",
            progress=70,
            detail=f"Удалено старых фрагментов: {old_chunks}; регионов: {old_regions}",
        )

        page_id_by_number: dict[int, int] = {}
        for item in page_items:
            page_row = Page(
                document_id=doc.id,
                organization_id=organization_id,
                project_id=project_id,
                page_number=int(item.get("page_number") or 0),
                width=int(item.get("width") or 0),
                height=int(item.get("height") or 0),
                image_path=str(item.get("image_path") or "") or None,
                text_raw=normalize_text(str(item.get("text") or "")) or None,
            )
            db.add(page_row)
            db.flush()
            page_id_by_number[page_row.page_number] = page_row.id

        ids: List[str] = []
        texts: List[str] = []
        metas: List[dict] = []
        chunks_created = 0
        for page_number, ch in _iter_chunks_with_pages(page_items):
            chunk_row = Chunk(
                document_id=doc.id,
                organization_id=organization_id,
                project_id=project_id,
                page_number=page_number,
                text=ch,
                region_id=None,
                embedding_id="__pending__",
            )
            db.add(chunk_row)
            db.flush()
            chunk_row.embedding_id = str(chunk_row.id)
            ids.append(str(chunk_row.id))
            texts.append(ch)
            metas.append(
                {
                    "chunk_id": chunk_row.id,
                    "document_id": doc.id,
                    "organization_id": organization_id or 0,
                    "project_id": project_id,
                    "page_number": page_number,
                    "page_id": page_id_by_number.get(page_number),
                    "filename": original_filename,
                    "region_id": None,
                    "region_type": "text_block",
                }
            )
            chunks_created += 1
        db.commit()

        if chunks_created and _should_index_embeddings():
            vectorstores.upsert_text(ids=ids, texts=texts, metadatas=metas)
        elif chunks_created:
            logger.info("Text embedding upsert skipped: RAG_INDEX_EMBEDDINGS is disabled")

        _set_document_state(
            db,
            doc,
            status="ready",
            stage="Индекс пересобран",
            progress=100,
            detail=f"Страниц: {len(page_items)} • фрагментов: {chunks_created}",
            finished=True,
        )
        logger.info("Reindexed document: id=%s pages=%s chunks=%s", doc.id, len(page_items), chunks_created)
        return int(doc.id), len(page_items), chunks_created, int(old_chunks)
    except Exception as exc:
        _remove_tree(tmp_page_dir)
        logger.exception("Document reindex failed: id=%s", doc.id)
        _set_document_state(
            db,
            doc,
            status="failed",
            stage="Ошибка пересборки индекса",
            progress=100,
            detail=f"{type(exc).__name__}: {exc}",
            finished=True,
        )
        raise


def ingest_document_text(
    *,
    db: Session,
    vectorstores: VectorStores,
    project_id: int,
    organization_id: int | None,
    file_type: str,
    stored_path: Path,
    original_filename: str,
) -> Tuple[int, int, int]:
    if file_type not in ("pdf", "docx"):
        raise ValueError("file_type must be pdf|docx")

    doc = Document(
        organization_id=organization_id,
        project_id=project_id,
        filename=original_filename,
        stored_path=str(stored_path),
        status="processing",
        stage="Документ принят в RAG",
        progress=5,
        detail="Ожидается извлечение содержимого",
        processing_started_at=datetime.utcnow(),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        page_dir = settings.PAGES_DIR / f"org_{organization_id or 0}" / str(project_id) / str(doc.id)
        page_dir.mkdir(parents=True, exist_ok=True)

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Извлечение страниц и текста",
            progress=20,
            detail="Подготавливаются страницы документа",
        )

        page_items = _extract_pdf_pages(stored_path, preview_dir=page_dir, filename=original_filename) if file_type == "pdf" else _extract_docx_pages(stored_path, preview_dir=page_dir, filename=original_filename)

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Подготовка страниц документа",
            progress=45,
            detail=f"Извлечено страниц: {len(page_items)}",
        )

        page_id_by_number: dict[int, int] = {}
        for item in page_items:
            page_row = Page(
                document_id=doc.id,
                organization_id=organization_id,
                project_id=project_id,
                page_number=int(item.get("page_number") or 0),
                width=int(item.get("width") or 0),
                height=int(item.get("height") or 0),
                image_path=str(item.get("image_path") or "") or None,
                text_raw=normalize_text(str(item.get("text") or "")) or None,
            )
            db.add(page_row)
            db.flush()
            page_id_by_number[page_row.page_number] = page_row.id

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Формирование текстовых фрагментов",
            progress=65,
            detail="Разбиваем документ на поисковые фрагменты",
        )

        ids: List[str] = []
        texts: List[str] = []
        metas: List[dict] = []
        chunks_created = 0

        for page_number, ch in _iter_chunks_with_pages(page_items):
            chunk_row = Chunk(
                document_id=doc.id,
                organization_id=organization_id,
                project_id=project_id,
                page_number=page_number,
                text=ch,
                region_id=None,
                embedding_id="__pending__",
            )
            db.add(chunk_row)
            db.flush()
            chunk_row.embedding_id = str(chunk_row.id)

            ids.append(str(chunk_row.id))
            texts.append(ch)
            metas.append(
                {
                    "chunk_id": chunk_row.id,
                    "document_id": doc.id,
                    "organization_id": organization_id or 0,
                    "project_id": project_id,
                    "page_number": page_number,
                    "page_id": page_id_by_number.get(page_number),
                    "filename": original_filename,
                    "region_id": None,
                    "region_type": "text_block",
                }
            )
            chunks_created += 1

        db.commit()

        _set_document_state(
            db,
            doc,
            status="processing",
            stage="Индексация в векторной базе",
            progress=85,
            detail=f"Подготовлено фрагментов: {chunks_created}",
        )

        if chunks_created and _should_index_embeddings():
            vectorstores.upsert_text(ids=ids, texts=texts, metadatas=metas)
        elif chunks_created:
            logger.info("Text embedding upsert skipped: RAG_INDEX_EMBEDDINGS is disabled")

        _set_document_state(
            db,
            doc,
            status="ready",
            stage="Документ готов",
            progress=100,
            detail=f"Страниц: {len(page_items)} • фрагментов: {chunks_created}",
            finished=True,
        )
        logger.info(
            "Ingested text document: file=%s project_id=%s doc_id=%s pages=%s chunks=%s",
            original_filename,
            project_id,
            doc.id,
            len(page_items),
            chunks_created,
        )
        return doc.id, len(page_items), chunks_created
    except Exception as exc:
        logger.exception("Ingestion failed for file=%s doc_id=%s: %s", original_filename, getattr(doc, "id", None), exc)
        try:
            _set_document_state(
                db,
                doc,
                status="failed",
                stage="Ошибка обработки",
                progress=max(1, int(getattr(doc, 'progress', 0) or 0)),
                detail=f"{type(exc).__name__}: {exc}",
                finished=True,
            )
        except Exception:
            logger.exception("Failed to persist failed document state for doc_id=%s", getattr(doc, 'id', None))
        try:
            _cleanup_failed_document(db, doc)
        except Exception:
            logger.exception("Failed to cleanup failed document doc_id=%s", getattr(doc, 'id', None))
        raise


def ingest_file(
    *,
    db: Session,
    vectorstores: VectorStores,
    project_id: int,
    organization_id: int | None,
    file_type: str,
    stored_path: Path,
    original_filename: str,
) -> Tuple[int, int]:
    doc_id, _pages_created, chunks_created = ingest_document_text(
        db=db,
        vectorstores=vectorstores,
        project_id=project_id,
        organization_id=organization_id,
        file_type=file_type,
        stored_path=stored_path,
        original_filename=original_filename,
    )
    return doc_id, chunks_created
